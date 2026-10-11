from __future__ import annotations

import html
import csv
import io
import re
import shutil
import subprocess
import tempfile
import urllib.error
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from .http_client import HttpClient
from .records import load_metadata
from .utils import flatten_unique, normalize_space, sha256_file, utc_now, write_json


PMC_OA_BASE = "https://www.ncbi.nlm.nih.gov/pmc/oai/oai.cgi"
PMC_EFETCH_BASE = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"


def sanitize_asset_name(value: str, fallback: str) -> str:
    value = html.unescape(value or "")
    value = urllib.parse.unquote(value)
    value = value.split("?", 1)[0].rstrip("/")
    name = Path(value).name or fallback
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return name or fallback


def absolutize_url(url: str, base: str) -> str:
    return urllib.parse.urljoin(base, html.unescape(url))


def extension_from_url(url: str) -> str:
    path = urllib.parse.urlparse(url).path.lower()
    return Path(path).suffix


def asset_type_from_text(text: str, url: str = "") -> str:
    combined = f"{text} {url}".lower()
    if "source data" in combined:
        return "source_data"
    if "supplementary" in combined or "supplement" in combined:
        if "table" in combined:
            return "supplementary_table"
        return "supplementary_material"
    if "table" in combined:
        return "table"
    if "data" in combined:
        return "data"
    return "asset"


def unique_assets(assets: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    seen = set()
    output = []
    for asset in assets:
        key = (asset.get("url", ""), asset.get("filename", ""))
        if key in seen:
            continue
        seen.add(key)
        output.append(asset)
    return output


def fetch_pmc_jats(http: HttpClient, pmcid: str) -> Tuple[Optional[str], str]:
    if not pmcid:
        return None, "no_pmcid"
    pmcid = pmcid if pmcid.upper().startswith("PMC") else f"PMC{pmcid}"
    url = (
        f"{PMC_OA_BASE}?verb=GetRecord&metadataPrefix=pmc"
        f"&identifier=oai:pubmedcentral.nih.gov:{pmcid[3:]}"
    )
    fallback_used = False
    try:
        text, _, _ = http.request_text(url)
    except urllib.error.HTTPError:
        fallback_url = f"{PMC_EFETCH_BASE}?db=pmc&id={pmcid[3:]}"
        text, _, _ = http.request_text(fallback_url)
        fallback_used = True
    if "<error" in text[:1000]:
        return None, "pmc_oai_unavailable"
    try:
        root = ET.fromstring(text)
    except ET.ParseError:
        return None, "pmc_oai_parse_error"
    article = None
    for elem in root.iter():
        if elem.tag.endswith("article"):
            article = elem
            break
    if article is None:
        return None, "pmc_article_not_found"
    status = "ok_efetch_fallback" if fallback_used else "ok"
    return ET.tostring(article, encoding="unicode"), status


def strip_namespace(root: ET.Element) -> None:
    for elem in root.iter():
        if "}" in elem.tag:
            elem.tag = elem.tag.split("}", 1)[1]


def elem_text(elem: Optional[ET.Element]) -> str:
    if elem is None:
        return ""
    return normalize_space("".join(elem.itertext()))


def save_article_text(record_dir: Path, article_xml: str) -> Dict[str, Any]:
    extracted_dir = record_dir / "extracted" / "text"
    extracted_dir.mkdir(parents=True, exist_ok=True)
    xml_path = extracted_dir / "pmc_article.xml"
    txt_path = extracted_dir / "article_text.txt"
    xml_path.write_text(article_xml, encoding="utf-8")

    root = ET.fromstring(article_xml)
    strip_namespace(root)
    title = elem_text(root.find(".//article-title"))
    abstract = elem_text(root.find(".//abstract"))
    body = elem_text(root.find(".//body"))
    back = elem_text(root.find(".//back"))
    text = "\n\n".join(part for part in [title, abstract, body, back] if part)
    txt_path.write_text(text + ("\n" if text else ""), encoding="utf-8")
    return {
        "xml_path": str(xml_path.relative_to(record_dir)),
        "text_path": str(txt_path.relative_to(record_dir)),
        "text_characters": len(text),
        "xml_sha256": sha256_file(xml_path),
        "text_sha256": sha256_file(txt_path),
    }


def parse_pmc_supplementary_assets(article_xml: str, base_url: str) -> List[Dict[str, Any]]:
    root = ET.fromstring(article_xml)
    strip_namespace(root)
    assets: List[Dict[str, Any]] = []

    for elem in root.findall(".//supplementary-material"):
        href = (
            elem.attrib.get("{http://www.w3.org/1999/xlink}href")
            or elem.attrib.get("href")
            or elem.attrib.get("xlink:href")
            or ""
        )
        if not href:
            continue
        label = elem_text(elem.find("label"))
        caption = elem_text(elem.find("caption")) or elem_text(elem)
        url = pmc_asset_url(href, base_url)
        fallback = f"supplementary_{len(assets) + 1}{extension_from_url(url) or '.dat'}"
        filename = sanitize_asset_name(url, fallback)
        assets.append(
            {
                "filename": filename,
                "url": url,
                "type": asset_type_from_text(f"{label} {caption}", url),
                "label": label,
                "caption": caption,
                "source": "pmc_jats",
                "local_path": "",
                "sha256": "",
                "downloaded_at": "",
                "download_status": "not_downloaded",
                "parse_status": "pending",
            }
        )

    for elem in root.findall(".//media"):
        href = (
            elem.attrib.get("{http://www.w3.org/1999/xlink}href")
            or elem.attrib.get("href")
            or elem.attrib.get("xlink:href")
            or ""
        )
        if not href:
            continue
        url = pmc_asset_url(href, base_url)
        fallback = f"media_{len(assets) + 1}{extension_from_url(url) or '.dat'}"
        filename = sanitize_asset_name(url, fallback)
        assets.append(
            {
                "filename": filename,
                "url": url,
                "type": asset_type_from_text(elem_text(elem), url),
                "label": "",
                "caption": elem_text(elem),
                "source": "pmc_jats",
                "local_path": "",
                "sha256": "",
                "downloaded_at": "",
                "download_status": "not_downloaded",
                "parse_status": "pending",
            }
        )

    return unique_assets(assets)


def pmc_asset_url(href: str, base_url: str) -> str:
    url = absolutize_url(href, base_url)
    parsed = urllib.parse.urlparse(url)
    if "/articles/instance/" in parsed.path and "/bin/" in parsed.path:
        return urllib.parse.urlunparse(parsed._replace(netloc="pmc.ncbi.nlm.nih.gov"))

    pmc_match = re.search(r"/(?:pmc/)?articles/(PMC)?(?P<pmcid>\d+)", parsed.path, re.IGNORECASE)
    if not pmc_match:
        pmc_match = re.search(r"PMC(?P<pmcid>\d+)", base_url, re.IGNORECASE)
    if not pmc_match:
        return url

    pmcid_number = pmc_match.group("pmcid")
    parsed = parsed._replace(netloc="pmc.ncbi.nlm.nih.gov")
    filename = Path(parsed.path).name
    if not filename:
        return url
    fixed_path = f"/articles/instance/{pmcid_number}/bin/{filename}"
    return urllib.parse.urlunparse(parsed._replace(path=fixed_path))


def extract_html_links(html_text: str, base_url: str, keywords: Iterable[str]) -> List[Dict[str, Any]]:
    keyword_values = [item.lower() for item in keywords]
    links: List[Dict[str, Any]] = []
    pattern = re.compile(
        r"<a\b(?P<attrs>[^>]*?)>(?P<body>.*?)</a>",
        flags=re.IGNORECASE | re.DOTALL,
    )
    href_pattern = re.compile(r'''href\s*=\s*["'](?P<href>[^"']+)["']''', re.IGNORECASE)
    tag_pattern = re.compile(r"<[^>]+>")
    for match in pattern.finditer(html_text):
        attrs = match.group("attrs")
        href_match = href_pattern.search(attrs)
        if not href_match:
            continue
        href = html.unescape(href_match.group("href"))
        body = normalize_space(tag_pattern.sub(" ", html.unescape(match.group("body"))))
        url = absolutize_url(href, base_url)
        combined = f"{body} {url}".lower()
        if not any(keyword in combined for keyword in keyword_values):
            continue
        fallback = f"publisher_asset_{len(links) + 1}{extension_from_url(url) or '.html'}"
        links.append(
            {
                "filename": sanitize_asset_name(url, fallback),
                "url": url,
                "type": asset_type_from_text(body, url),
                "label": "",
                "caption": body,
                "source": "publisher_html",
                "local_path": "",
                "sha256": "",
                "downloaded_at": "",
                "download_status": "not_downloaded",
                "parse_status": "pending",
            }
        )
    return unique_assets(links)


def normalize_elsevier_pii(value: str) -> str:
    value = html.unescape(value or "")
    match = re.search(r"S(?P<digits>\d{16})", value)
    if match:
        return f"S{match.group('digits')}"
    match = re.search(
        r"S(?P<issn>\d{4})-(?P<issn2>\d{4})\((?P<year>\d{2})\)(?P<num>\d{5})-(?P<check>\d)",
        value,
    )
    if match:
        return (
            f"S{match.group('issn')}{match.group('issn2')}"
            f"{match.group('year')}{match.group('num')}{match.group('check')}"
        )
    return ""


def extract_elsevier_pii(text: str) -> str:
    return normalize_elsevier_pii(text)


def elsevier_cdn_asset(
    pii: str,
    index: int,
    extension: str,
    content_type: str,
    content_length: int,
) -> Dict[str, Any]:
    url = f"https://ars.els-cdn.com/content/image/1-s2.0-{pii}-mmc{index}.{extension}"
    filename = f"elsevier_{pii}_mmc{index}.{extension}"
    return {
        "filename": filename,
        "url": url,
        "type": asset_type_from_text(f"mmc{index} {extension}", url),
        "label": f"mmc{index}",
        "caption": f"Elsevier supplementary material mmc{index}",
        "source": "elsevier_cdn_guess",
        "local_path": "",
        "sha256": "",
        "downloaded_at": "",
        "download_status": "not_downloaded",
        "parse_status": "pending",
        "content_type": content_type,
        "content_length": content_length,
        "pii": pii,
    }


def probe_elsevier_cdn_assets(
    http: HttpClient,
    pii: str,
    extensions: Iterable[str],
    max_mmc: int = 20,
) -> List[Dict[str, Any]]:
    if not pii:
        return []
    assets: List[Dict[str, Any]] = []
    for index in range(1, max_mmc + 1):
        found_for_index = False
        for extension in extensions:
            extension = extension.lower().lstrip(".")
            url = f"https://ars.els-cdn.com/content/image/1-s2.0-{pii}-mmc{index}.{extension}"
            try:
                status, _, content_type, content_length = http.head(url)
            except Exception:
                continue
            if status == 200 and content_length > 0:
                assets.append(
                    elsevier_cdn_asset(
                        pii=pii,
                        index=index,
                        extension=extension,
                        content_type=content_type,
                        content_length=content_length,
                    )
                )
                found_for_index = True
                break
        if not found_for_index and index > 4:
            break
    return assets


def needs_elsevier_cdn_fallback(
    existing_assets: Iterable[Dict[str, Any]],
    incoming_assets: Iterable[Dict[str, Any]],
) -> bool:
    existing_list = list(existing_assets)
    incoming_list = list(incoming_assets)
    blocked_statuses = {"blocked_challenge", "invalid_content", "failed"}
    if any(asset.get("download_status") in blocked_statuses for asset in existing_list):
        return True
    if not existing_list and not incoming_list:
        return True
    return False


def load_manifest(record_dir: Path) -> List[Dict[str, Any]]:
    path = record_dir / "assets" / "assets_manifest.json"
    if not path.exists() or path.stat().st_size == 0:
        return []
    import json

    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_manifest(record_dir: Path, assets: List[Dict[str, Any]]) -> None:
    write_json(record_dir / "assets" / "assets_manifest.json", assets)


def asset_extension(asset: Dict[str, Any]) -> str:
    name = asset.get("filename") or asset.get("local_path") or asset.get("url", "")
    suffix = Path(urllib.parse.urlparse(name).path).suffix.lower()
    if suffix:
        return suffix
    return extension_from_url(asset.get("url", ""))


def is_download_candidate(asset: Dict[str, Any], extensions: Iterable[str]) -> bool:
    ext = asset_extension(asset)
    allowed = {item.lower() for item in extensions}
    return ext in allowed


def validate_downloaded_asset(
    asset: Dict[str, Any], path: Path, content_type: str
) -> Tuple[bool, str, str]:
    ext = asset_extension(asset)
    content_type_lower = (content_type or "").lower()
    with path.open("rb") as handle:
        head_bytes = handle.read(8192)
    head_text = head_bytes.decode("utf-8", errors="ignore").lower()
    html_markers = (
        "<!doctype html",
        "<html",
        "recaptcha",
        "checking your browser",
        "pow_challenge",
        "cloudpmc-viewer-pow",
        "captcha",
    )
    looks_html = "text/html" in content_type_lower or any(
        marker in head_text for marker in html_markers
    )
    if looks_html:
        if "recaptcha" in head_text or "checking your browser" in head_text:
            return False, "blocked_challenge", "HTML reCAPTCHA/browser-check page was downloaded instead of the asset"
        if "pow_challenge" in head_text or "cloudpmc-viewer-pow" in head_text:
            return False, "blocked_challenge", "PMC proof-of-work page was downloaded instead of the asset"
        return False, "invalid_content", "HTML page was downloaded instead of the asset"

    is_zip_office = head_bytes.startswith(b"PK")
    is_ole_office = head_bytes.startswith(b"\xd0\xcf\x11\xe0")
    if ext in {".xlsx", ".xlsm", ".docx"} and not (
        is_zip_office or (ext == ".xlsx" and is_ole_office)
    ):
        return False, "invalid_content", f"{ext} asset is not a zip-based Office file"
    if ext == ".xls" and not (is_ole_office or is_zip_office):
        data_lines = [
            line.strip()
            for line in head_text.splitlines()
            if line.strip() and not line.lstrip().startswith("#")
        ][:20]
        looks_delimited = any(
            sum(line.count(delimiter) > 0 for line in data_lines) >= 2
            for delimiter in ("\t", ",", ";")
        )
        if not looks_delimited:
            return (
                False,
                "invalid_content",
                ".xls asset is not an Excel binary/zip or delimited text file",
            )
    return True, "", ""


def download_assets_for_record(
    record_dir: Path,
    http: HttpClient,
    extensions: Iterable[str],
    overwrite: bool = False,
    elsevier_api_key: str = "",
    elsevier_inst_token: str = "",
) -> Dict[str, Any]:
    assets = load_manifest(record_dir)
    assets_dir = record_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    stats = {
        "asset_count": len(assets),
        "candidate_count": 0,
        "downloaded": 0,
        "skipped": 0,
        "skipped_extension": 0,
        "failed": 0,
        "errors": [],
    }

    used_names = set()
    for asset in assets:
        if asset.get("source") == "pmc_jats" and asset.get("url"):
            asset["url"] = pmc_asset_url(asset["url"], asset["url"])
        filename = sanitize_asset_name(asset.get("filename", ""), "asset.dat")
        if filename in used_names:
            stem = Path(filename).stem
            suffix = Path(filename).suffix
            filename = f"{stem}_{len(used_names) + 1}{suffix}"
        used_names.add(filename)
        asset["filename"] = filename

        if not asset.get("url"):
            asset["download_status"] = "missing_url"
            stats["skipped"] += 1
            continue
        if not is_download_candidate(asset, extensions):
            stats["skipped_extension"] += 1
            continue
        stats["candidate_count"] += 1

        local_path = assets_dir / filename
        if local_path.exists() and not overwrite:
            asset["local_path"] = str(local_path.relative_to(record_dir))
            asset["sha256"] = sha256_file(local_path)
            valid, status, message = validate_downloaded_asset(
                asset, local_path, asset.get("content_type", "")
            )
            if not valid:
                asset["download_status"] = status
                asset["download_error"] = message
            else:
                asset["download_status"] = "exists"
                asset.pop("download_error", None)
                stats["skipped"] += 1
                continue

        try:
            request_headers = None
            if asset.get("api_auth") == "elsevier":
                if not elsevier_api_key:
                    raise PermissionError("ELSEVIER_API_KEY is required for this asset")
                request_headers = {
                    "Accept": "*/*",
                    "X-ELS-APIKey": elsevier_api_key,
                }
                if elsevier_inst_token:
                    request_headers["X-ELS-Insttoken"] = elsevier_inst_token
            final_url, content_type = http.download(
                asset["url"],
                local_path,
                headers=request_headers,
            )
            asset["url_final"] = final_url
            asset["content_type"] = content_type
            asset["local_path"] = str(local_path.relative_to(record_dir))
            asset["sha256"] = sha256_file(local_path)
            asset["downloaded_at"] = utc_now()
            valid, status, message = validate_downloaded_asset(asset, local_path, content_type)
            if not valid:
                asset["download_status"] = status
                asset["download_error"] = message
                stats["failed"] += 1
                stats["errors"].append(f"{filename}: {message}")
                continue
            asset["download_status"] = "downloaded"
            asset.pop("download_error", None)
            stats["downloaded"] += 1
        except Exception as exc:  # noqa: BLE001 - per-asset download should continue.
            asset["download_status"] = "failed"
            asset["download_error"] = f"{type(exc).__name__}: {exc}"
            stats["failed"] += 1
            stats["errors"].append(f"{filename}: {asset['download_error']}")

    write_manifest(record_dir, assets)
    return stats


def safe_table_name(value: str, fallback: str) -> str:
    value = normalize_space(value or "")
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return value[:120].rstrip("._") or fallback


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def write_csv_table(path: Path, rows: Iterable[Iterable[Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            writer.writerow([cell_to_text(cell) for cell in row])
            count += 1
    return count


def repair_xlsx_archive(local_path: Path, error: Exception, work_dir: Path) -> Path:
    repaired_path = work_dir / "repaired.xlsx"
    if isinstance(error, zipfile.BadZipFile):
        unpacked_dir = work_dir / "unpacked"
        unpacked_dir.mkdir()
        subprocess.run(
            ["unzip", "-q", str(local_path), "-d", str(unpacked_dir)],
            check=True,
            capture_output=True,
        )
        archive_path = shutil.make_archive(
            str(work_dir / "repaired"), "zip", root_dir=unpacked_dir
        )
        Path(archive_path).replace(repaired_path)
        return repaired_path

    with zipfile.ZipFile(local_path) as source, zipfile.ZipFile(
        repaired_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for member in source.infolist():
            if member.filename == "docProps/custom.xml":
                continue
            data = source.read(member.filename)
            if member.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r'<Override[^>]+PartName="/docProps/custom\.xml"[^>]*/>',
                    "",
                    text,
                )
                data = text.encode("utf-8")
            target.writestr(member, data)
    return repaired_path


def extract_xlsx_tables(record_dir: Path, asset: Dict[str, Any]) -> List[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on environment.
        raise RuntimeError("openpyxl is required to parse xlsx/xlsm files") from exc

    local_path = record_dir / asset.get("local_path", "")
    with local_path.open("rb") as handle:
        signature = handle.read(4)
    if signature.startswith(b"\xd0\xcf\x11\xe0"):
        return extract_xls_tables(record_dir, asset)

    workbook_source: Any = local_path
    source_handle = None
    if local_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        source_handle = local_path.open("rb")
        workbook_source = source_handle
    repair_dir = None
    try:
        workbook = openpyxl.load_workbook(
            workbook_source, read_only=True, data_only=True
        )
    except (TypeError, zipfile.BadZipFile) as exc:
        if source_handle is not None:
            source_handle.close()
            source_handle = None
        repair_dir = tempfile.TemporaryDirectory(prefix="marker_xlsx_repair_")
        repaired_path = repair_xlsx_archive(
            local_path, exc, Path(repair_dir.name)
        )
        workbook = openpyxl.load_workbook(
            repaired_path, read_only=True, data_only=True
        )
    outputs: List[Dict[str, Any]] = []
    base = safe_table_name(Path(asset.get("filename", local_path.name)).stem, "workbook")
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        sheet_name = safe_table_name(worksheet.title, f"sheet_{sheet_index}")
        output_path = (
            record_dir
            / "extracted"
            / "tables"
            / f"{base}__{sheet_index:02d}_{sheet_name}.csv"
        )
        row_count = write_csv_table(output_path, worksheet.iter_rows(values_only=True))
        outputs.append(
            {
                "source_asset": asset.get("filename", ""),
                "source_local_path": asset.get("local_path", ""),
                "table_type": "xlsx_sheet",
                "sheet_name": worksheet.title,
                "output_path": str(output_path.relative_to(record_dir)),
                "rows": row_count,
                "sha256": sha256_file(output_path),
            }
        )
    workbook.close()
    if source_handle is not None:
        source_handle.close()
    if repair_dir is not None:
        repair_dir.cleanup()
    return outputs


def is_xlsx_archive(local_path: Path) -> bool:
    try:
        with zipfile.ZipFile(local_path) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        return False
    return "[Content_Types].xml" in names and "xl/workbook.xml" in names


def text_encoding(local_path: Path) -> str:
    with local_path.open("rb") as handle:
        signature = handle.read(4)
    if signature.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    return "utf-8-sig"


def sniff_delimiter(sample: str, default: str = ",") -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        counts = {delimiter: sample.count(delimiter) for delimiter in ("\t", ",", ";", "|")}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        return delimiter if count else default


def extract_mislabeled_delimited_table(
    record_dir: Path, asset: Dict[str, Any]
) -> List[Dict[str, Any]]:
    local_path = record_dir / asset.get("local_path", "")
    encoding = text_encoding(local_path)
    with local_path.open("r", encoding=encoding, errors="replace", newline="") as handle:
        sample = handle.read(65536)
    delimiter = sniff_delimiter(sample, default="\t")
    output_path = (
        record_dir
        / "extracted"
        / "tables"
        / f"{safe_table_name(local_path.stem, 'table')}.csv"
    )
    with local_path.open("r", encoding=encoding, errors="replace", newline="") as handle:
        row_count = write_csv_table(output_path, csv.reader(handle, delimiter=delimiter))
    return [
        {
            "source_asset": asset.get("filename", ""),
            "source_local_path": asset.get("local_path", ""),
            "table_type": "mislabeled_delimited_text",
            "sheet_name": "",
            "output_path": str(output_path.relative_to(record_dir)),
            "rows": row_count,
            "sha256": sha256_file(output_path),
        }
    ]


def extract_zip_wrapped_tables(
    record_dir: Path, asset: Dict[str, Any]
) -> List[Dict[str, Any]]:
    local_path = record_dir / asset.get("local_path", "")
    outputs: List[Dict[str, Any]] = []
    outer_base = safe_table_name(local_path.stem, "archive")
    with zipfile.ZipFile(local_path) as archive:
        members = [
            member
            for member in archive.infolist()
            if not member.is_dir()
            and not member.filename.startswith("__MACOSX/")
            and not Path(member.filename).name.startswith("._")
            and Path(member.filename).suffix.lower() in {".csv", ".tsv"}
        ]
        if not members:
            raise RuntimeError("zip-wrapped table asset contains no CSV or TSV file")
        for member_index, member in enumerate(members, start=1):
            inner_name = Path(member.filename).name
            inner_base = safe_table_name(Path(inner_name).stem, f"table_{member_index}")
            output_path = (
                record_dir
                / "extracted"
                / "tables"
                / f"{outer_base}__{member_index:02d}_{inner_base}.csv"
            )
            with archive.open(member) as raw_handle:
                sample_bytes = raw_handle.read(65536)
            encoding = "utf-16" if sample_bytes.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            sample = sample_bytes.decode(encoding, errors="replace")
            default = "\t" if Path(inner_name).suffix.lower() == ".tsv" else ","
            delimiter = sniff_delimiter(sample, default=default)
            with archive.open(member) as raw_handle, io.TextIOWrapper(
                raw_handle, encoding=encoding, errors="replace", newline=""
            ) as text_handle:
                row_count = write_csv_table(
                    output_path, csv.reader(text_handle, delimiter=delimiter)
                )
            outputs.append(
                {
                    "source_asset": asset.get("filename", ""),
                    "source_local_path": asset.get("local_path", ""),
                    "table_type": "zip_delimited_text",
                    "sheet_name": inner_name,
                    "output_path": str(output_path.relative_to(record_dir)),
                    "rows": row_count,
                    "sha256": sha256_file(output_path),
                }
            )
    return outputs


def extract_xls_tables(record_dir: Path, asset: Dict[str, Any]) -> List[Dict[str, Any]]:
    local_path = record_dir / asset.get("local_path", "")
    with local_path.open("rb") as handle:
        signature = handle.read(4)
    if signature.startswith(b"PK"):
        if is_xlsx_archive(local_path):
            return extract_xlsx_tables(record_dir, asset)
        return extract_zip_wrapped_tables(record_dir, asset)
    if not signature.startswith(b"\xd0\xcf\x11\xe0"):
        return extract_mislabeled_delimited_table(record_dir, asset)

    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - depends on environment.
        raise RuntimeError("xlrd is required to parse legacy .xls files") from exc

    workbook = xlrd.open_workbook(str(local_path), on_demand=True)
    outputs: List[Dict[str, Any]] = []
    base = safe_table_name(Path(asset.get("filename", local_path.name)).stem, "workbook")
    for sheet_index in range(workbook.nsheets):
        worksheet = workbook.sheet_by_index(sheet_index)
        sheet_name = safe_table_name(worksheet.name, f"sheet_{sheet_index + 1}")
        output_path = (
            record_dir
            / "extracted"
            / "tables"
            / f"{base}__{sheet_index + 1:02d}_{sheet_name}.csv"
        )
        row_count = write_csv_table(
            output_path,
            (worksheet.row_values(row_index) for row_index in range(worksheet.nrows)),
        )
        outputs.append(
            {
                "source_asset": asset.get("filename", ""),
                "source_local_path": asset.get("local_path", ""),
                "table_type": "xls_sheet",
                "sheet_name": worksheet.name,
                "output_path": str(output_path.relative_to(record_dir)),
                "rows": row_count,
                "sha256": sha256_file(output_path),
            }
        )
    workbook.release_resources()
    return outputs


def extract_csv_or_tsv_table(record_dir: Path, asset: Dict[str, Any]) -> List[Dict[str, Any]]:
    local_path = record_dir / asset.get("local_path", "")
    ext = local_path.suffix.lower()
    if ext == ".csv":
        return [
            {
                "source_asset": asset.get("filename", ""),
                "source_local_path": asset.get("local_path", ""),
                "table_type": "csv_asset",
                "sheet_name": "",
                "output_path": asset.get("local_path", ""),
                "rows": None,
                "row_count_status": "not_counted",
                "sha256": asset.get("sha256") or sha256_file(local_path),
            }
        ]

    delimiter = "\t" if ext == ".tsv" else ","
    output_path = (
        record_dir
        / "extracted"
        / "tables"
        / f"{safe_table_name(local_path.stem, 'table')}.csv"
    )
    with local_path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        row_count = write_csv_table(output_path, reader)
    return [
        {
            "source_asset": asset.get("filename", ""),
            "source_local_path": asset.get("local_path", ""),
            "table_type": "delimited_text",
            "sheet_name": "",
            "output_path": str(output_path.relative_to(record_dir)),
            "rows": row_count,
            "sha256": sha256_file(output_path),
        }
    ]


def extract_docx_tables(record_dir: Path, asset: Dict[str, Any]) -> List[Dict[str, Any]]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on environment.
        raise RuntimeError("python-docx is required to parse docx files") from exc

    local_path = record_dir / asset.get("local_path", "")
    document = docx.Document(str(local_path))
    outputs: List[Dict[str, Any]] = []
    base = safe_table_name(local_path.stem, "docx")
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        output_path = (
            record_dir
            / "extracted"
            / "tables"
            / f"{base}__table_{table_index:02d}.csv"
        )
        row_count = write_csv_table(output_path, rows)
        outputs.append(
            {
                "source_asset": asset.get("filename", ""),
                "source_local_path": asset.get("local_path", ""),
                "table_type": "docx_table",
                "sheet_name": f"table_{table_index}",
                "output_path": str(output_path.relative_to(record_dir)),
                "rows": row_count,
                "sha256": sha256_file(output_path),
            }
        )
    return outputs


def sync_local_table_assets(
    record_dir: Path, assets: List[Dict[str, Any]]
) -> int:
    assets_dir = record_dir / "assets"
    if not assets_dir.exists():
        return 0
    supported = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".docx"}
    tracked_paths = {asset.get("local_path", "") for asset in assets}
    by_filename: Dict[str, List[Dict[str, Any]]] = {}
    for asset in assets:
        by_filename.setdefault(asset.get("filename", ""), []).append(asset)

    synced = 0
    for local_path in sorted(assets_dir.iterdir()):
        if not local_path.is_file() or local_path.suffix.lower() not in supported:
            continue
        relative_path = str(local_path.relative_to(record_dir))
        if relative_path in tracked_paths:
            continue
        candidates = [
            asset
            for asset in by_filename.get(local_path.name, [])
            if not asset.get("local_path")
        ]
        if candidates:
            asset = candidates[0]
            asset["local_path"] = relative_path
            asset["sha256"] = sha256_file(local_path)
            asset["download_status"] = "exists"
            asset["parse_status"] = "pending"
        else:
            asset = {
                "filename": local_path.name,
                "url": "",
                "type": "local_table_asset",
                "label": "",
                "caption": "Locally recovered table asset",
                "source": "local_asset_sync",
                "local_path": relative_path,
                "sha256": sha256_file(local_path),
                "download_status": "exists",
                "parse_status": "pending",
            }
            assets.append(asset)
            by_filename.setdefault(local_path.name, []).append(asset)
        tracked_paths.add(relative_path)
        synced += 1
    return synced


def extract_tables_for_record(record_dir: Path) -> Dict[str, Any]:
    assets = load_manifest(record_dir)
    synced_assets = sync_local_table_assets(record_dir, assets)
    tables_dir = record_dir / "extracted" / "tables"
    tables_manifest_path = tables_dir / "tables_manifest.json"
    if tables_dir.exists():
        for output_path in tables_dir.glob("*.csv"):
            output_path.unlink()
    outputs: List[Dict[str, Any]] = []
    stats = {
        "asset_count": len(assets),
        "synced_assets": synced_assets,
        "parsed_assets": 0,
        "tables": 0,
        "duplicates": 0,
        "skipped": 0,
        "failed": 0,
        "errors": [],
    }

    seen_checksums: Dict[str, str] = {}
    for asset in assets:
        local_path_value = asset.get("local_path", "")
        if not local_path_value:
            asset["parse_status"] = "not_downloaded"
            asset.pop("parse_error", None)
            stats["skipped"] += 1
            continue
        local_path = record_dir / local_path_value
        if not local_path.exists():
            asset["parse_status"] = "missing_local_file"
            asset.pop("parse_error", None)
            stats["skipped"] += 1
            continue

        valid, status, message = validate_downloaded_asset(
            asset, local_path, asset.get("content_type", "")
        )
        if not valid:
            asset["parse_status"] = f"skipped_{status}"
            asset["parse_error"] = message
            asset.pop("extracted_tables", None)
            stats["skipped"] += 1
            continue
        if asset.get("download_status") not in {"downloaded", "exists"}:
            asset["download_status"] = "exists"
            asset.pop("download_error", None)

        checksum = asset.get("sha256") or sha256_file(local_path)
        asset["sha256"] = checksum
        if checksum in seen_checksums:
            asset["parse_status"] = "duplicate_content"
            asset["duplicate_of"] = seen_checksums[checksum]
            asset.pop("extracted_tables", None)
            asset.pop("parse_error", None)
            stats["duplicates"] += 1
            stats["skipped"] += 1
            continue
        seen_checksums[checksum] = asset.get("filename", local_path.name)
        asset.pop("duplicate_of", None)

        ext = local_path.suffix.lower()
        try:
            extracted: List[Dict[str, Any]]
            if ext in {".xlsx", ".xlsm"}:
                extracted = extract_xlsx_tables(record_dir, asset)
            elif ext == ".xls":
                extracted = extract_xls_tables(record_dir, asset)
            elif ext in {".csv", ".tsv"}:
                with local_path.open("rb") as handle:
                    signature = handle.read(4)
                if signature.startswith(b"PK"):
                    extracted = (
                        extract_xlsx_tables(record_dir, asset)
                        if is_xlsx_archive(local_path)
                        else extract_zip_wrapped_tables(record_dir, asset)
                    )
                else:
                    extracted = extract_csv_or_tsv_table(record_dir, asset)
            elif ext == ".docx":
                extracted = extract_docx_tables(record_dir, asset)
            else:
                asset["parse_status"] = "unsupported_extension"
                asset.pop("parse_error", None)
                stats["skipped"] += 1
                continue

            asset["parse_status"] = "parsed" if extracted else "no_tables_found"
            asset["parsed_at"] = utc_now()
            asset["extracted_tables"] = extracted
            asset.pop("parse_error", None)
            outputs.extend(extracted)
            stats["parsed_assets"] += 1
            stats["tables"] += len(extracted)
        except Exception as exc:  # noqa: BLE001 - per-asset parsing should continue.
            asset["parse_status"] = "failed"
            asset["parse_error"] = f"{type(exc).__name__}: {exc}"
            stats["failed"] += 1
            stats["errors"].append(f"{asset.get('filename', '')}: {asset['parse_error']}")

    write_manifest(record_dir, assets)
    write_json(tables_manifest_path, outputs)
    return stats


def merge_assets(existing: List[Dict[str, Any]], incoming: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    merged: List[Dict[str, Any]] = []
    by_key: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for asset in existing + incoming:
        key = (asset.get("url", ""), asset.get("filename", ""))
        if key in by_key:
            current = by_key[key]
            for field, value in asset.items():
                if value not in ("", None, [], {}) and not current.get(field):
                    current[field] = value
        else:
            copied = dict(asset)
            by_key[key] = copied
            merged.append(copied)
    return merged


def probe_record(
    record_dir: Path,
    metadata: Dict[str, Any],
    http: HttpClient,
    config: Dict[str, Any],
    fetch_publisher: bool = True,
    save_text: bool = True,
    probe_elsevier_cdn: bool = True,
    publisher_if_no_assets: bool = False,
    fetch_elsevier_publisher: bool = True,
) -> Dict[str, Any]:
    pmcid = metadata.get("pmcid", "")
    probe: Dict[str, Any] = {
        "probed_at": utc_now(),
        "pmc_status": "not_attempted",
        "publisher_status": "not_attempted",
        "article_text": {},
        "asset_count": 0,
        "new_asset_count": 0,
        "errors": [],
    }

    incoming_assets: List[Dict[str, Any]] = []
    article_xml = None
    try:
        article_xml, status = fetch_pmc_jats(http, pmcid)
        probe["pmc_status"] = status
    except Exception as exc:  # noqa: BLE001 - publisher fallback may still recover assets.
        probe["pmc_status"] = "error"
        probe["errors"].append(f"pmc_fetch: {type(exc).__name__}: {exc}")
    elsevier_pii = ""
    if article_xml:
        try:
            if save_text:
                text_info = save_article_text(record_dir, article_xml)
                probe["article_text"] = text_info
            base_url = f"https://www.ncbi.nlm.nih.gov/pmc/articles/{pmcid}/"
            incoming_assets.extend(parse_pmc_supplementary_assets(article_xml, base_url))
            elsevier_pii = extract_elsevier_pii(article_xml)
        except Exception as exc:  # noqa: BLE001 - keep per-record probing robust.
            probe["errors"].append(f"pmc_parse: {type(exc).__name__}: {exc}")
    else:
        cached_xml = record_dir / "extracted" / "text" / "pmc_article.xml"
        if cached_xml.exists():
            try:
                elsevier_pii = extract_elsevier_pii(
                    cached_xml.read_text(encoding="utf-8", errors="replace")
                )
            except Exception as exc:  # noqa: BLE001 - cache is only a fallback hint.
                probe["errors"].append(f"cached_pmc_xml: {type(exc).__name__}: {exc}")

    is_elsevier = metadata.get("doi", "").lower().startswith("10.1016/")
    should_fetch_publisher = (
        fetch_publisher
        and metadata.get("url")
        and (fetch_elsevier_publisher or not is_elsevier)
        and (not publisher_if_no_assets or not incoming_assets)
    )
    if should_fetch_publisher:
        try:
            html_text, final_url, _ = http.request_text(metadata["url"])
            elsevier_pii = elsevier_pii or extract_elsevier_pii(html_text + " " + final_url)
            publisher_links = extract_html_links(
                html_text,
                final_url,
                config.get("asset_link_keywords", []),
            )
            incoming_assets.extend(publisher_links)
            probe["publisher_status"] = "ok"
        except Exception as exc:  # noqa: BLE001 - publisher pages vary; keep going.
            probe["publisher_status"] = "error"
            probe["errors"].append(f"publisher_html: {type(exc).__name__}: {exc}")

    existing_assets = load_manifest(record_dir)

    if (
        probe_elsevier_cdn
        and fetch_publisher
        and is_elsevier
        and needs_elsevier_cdn_fallback(existing_assets, incoming_assets)
    ):
        try:
            cdn_assets = probe_elsevier_cdn_assets(
                http=http,
                pii=elsevier_pii,
                extensions=config.get("asset_extensions", []),
            )
            incoming_assets.extend(cdn_assets)
            if cdn_assets:
                probe["publisher_status"] = (
                    "ok_elsevier_cdn"
                    if probe["publisher_status"] in {"not_attempted", "error"}
                    else f"{probe['publisher_status']}+elsevier_cdn"
                )
                probe["elsevier_pii"] = elsevier_pii
                probe["elsevier_cdn_asset_count"] = len(cdn_assets)
        except Exception as exc:  # noqa: BLE001 - fallback probing must not break PMC path.
            probe["errors"].append(f"elsevier_cdn: {type(exc).__name__}: {exc}")

    merged_assets = merge_assets(existing_assets, unique_assets(incoming_assets))
    write_manifest(record_dir, merged_assets)

    metadata["assets"] = merged_assets
    metadata.setdefault("content_probe", {})
    metadata["content_probe"] = probe
    metadata["updated_at"] = utc_now()
    probe["asset_count"] = len(merged_assets)
    probe["new_asset_count"] = max(0, len(merged_assets) - len(existing_assets))
    return probe
